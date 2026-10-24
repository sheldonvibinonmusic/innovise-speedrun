"""Pytest fixtures and test environment configuration.

Provides isolated temporary SQLite databases, custom upload directories,
in-memory Qdrant vector stores, deterministic embedders, and synthetic PDF fixtures.
"""

from pathlib import Path
from typing import Generator
import pymupdf
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from app.config import settings
from app.db.models import Base
from app.db.session import get_db
from app.main import create_app
from app.retrieval.embeddings import DeterministicEmbedder
from app.retrieval.service import retrieval_service
from app.retrieval.vector_store import QdrantVectorStore


@pytest.fixture(scope="session")
def valid_pdf_bytes() -> bytes:
    """Minimal valid PDF byte sequence starting with %PDF- header."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Standard sample document for minimal validation.")
    content = doc.tobytes()
    doc.close()
    return content


@pytest.fixture(scope="session")
def financial_10k_pdf_bytes() -> bytes:
    """Synthetic standard Form 10-K with MD&A narrative and structured financial table."""
    doc = pymupdf.open()
    # Page 1: Overview and MD&A
    p1 = doc.new_page()
    p1.insert_textbox(
        pymupdf.Rect(50, 50, 550, 750),
        "Microsoft Corporation - Annual Report on Form 10-K for Fiscal Year 2024\n\n"
        "Item 7. Management's Discussion and Analysis of Financial Condition\n\n"
        "Our cloud services continued strong operational momentum throughout FY24. "
        "Total revenue increased 16% year-over-year driven by Intelligent Cloud expansion. "
        "Operating margins remained highly profitable at 44.6%, reflecting disciplined capital allocation. "
        "Operating cash flow reached record highs, providing ample liquidity for cloud infrastructure investments.\n",
    )

    # Page 2: Financial Statements Table
    p2 = doc.new_page()
    p2.insert_textbox(
        pymupdf.Rect(50, 50, 550, 750),
        "CONSOLIDATED STATEMENTS OF INCOME (in millions)\n\n"
        "Line Item | 2024 | 2023\n"
        "Total revenue | 245120 | 211915\n"
        "Operating income | 109433 | 88523\n"
        "Net income | 88136 | 72361\n"
        "Operating cash flow | 118548 | 87582\n"
        "Cash and cash equivalents | 75544 | 34704\n"
        "Total assets | 512163 | 411976\n"
        "Total debt | 44922 | 47204\n",
    )
    content = doc.tobytes()
    doc.close()
    return content


@pytest.fixture(scope="session")
def conflicting_10k_pdf_bytes() -> bytes:
    """Synthetic conflicting disclosure reporting revenue decline and margin compression."""
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_textbox(
        pymupdf.Rect(50, 50, 550, 750),
        "Microsoft Corporation - Restated Fiscal Year 2024 Analysis\n\n"
        "Item 1A. Risk Factors & Restated Statements\n\n"
        "The company experienced sharp revenue decline and significant margin compression. "
        "Headwinds in commercial hardware and customer slowdown led to substantial operating losses. "
        "Total revenue decreased drastically due to severe market contraction.\n\n"
        "CONSOLIDATED STATEMENTS OF OPERATIONS (in millions)\n\n"
        "Line Item | 2024 | 2023\n"
        "Total revenue | 185000 | 211915\n"
        "Net income | 42000 | 72361\n",
    )
    content = doc.tobytes()
    doc.close()
    return content


@pytest.fixture(scope="session")
def blank_page_pdf_bytes() -> bytes:
    """PDF containing a blank page to trigger extraction warnings."""
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((50, 50), "First page contains digital text.")
    # Page 2 is empty
    doc.new_page()
    content = doc.tobytes()
    doc.close()
    return content


@pytest.fixture(scope="session")
def financial_presentation_bytes() -> bytes:
    """Synthetic investor presentation (.pptx) with financial slides, tables, and notes."""
    import io
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "FY24 Investor Strategic Overview"
    slide1.placeholders[1].text = "Cloud segment revenue accelerated across global regions."
    slide1.notes_slide.notes_text_frame.text = "Executive note: emphasize margin durability."

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Segment Operating Performance"
    table_shape = slide2.shapes.add_table(3, 2, Inches(1), Inches(1.5), Inches(5), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Segment"
    table.cell(0, 1).text = "Operating Revenue"
    table.cell(1, 0).text = "Cloud Services"
    table.cell(1, 1).text = "$105,360M"
    table.cell(2, 0).text = "Productivity"
    table.cell(2, 1).text = "$69,200M"

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()


@pytest.fixture
def temp_env(tmp_path: Path):
    """Isolate database, upload, Qdrant in-memory, and BM25 index per test."""
    test_db_file = tmp_path / "test_financial.db"
    test_db_url = f"sqlite:///{test_db_file.as_posix()}"
    test_upload_dir = tmp_path / "uploads"
    test_upload_dir.mkdir(parents=True, exist_ok=True)

    # Save original settings
    original_db_url = settings.database_url
    original_upload_dir = settings.upload_dir

    # Apply test settings
    settings.database_url = test_db_url
    settings.upload_dir = test_upload_dir

    # Create isolated engine and tables
    engine = create_engine(test_db_url, connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

    import app.db.session as session_module
    old_engine = session_module._engine
    old_sf = session_module._session_factory
    session_module._engine = engine
    session_module._session_factory = TestingSessionLocal

    # Configure isolated in-memory Qdrant and clean BM25
    isolated_vector_store = QdrantVectorStore(in_memory=True)
    retrieval_service.vector_store = isolated_vector_store
    retrieval_service.embedding_service = DeterministicEmbedder()
    retrieval_service.bm25_index.corpus_metadata = []
    retrieval_service.bm25_index.chunk_ids = []
    retrieval_service.bm25_index.bm25 = None
    retrieval_service.reranker.enabled = False

    app = create_app()

    def override_get_db() -> Generator[Session, None, None]:
        session = TestingSessionLocal()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db

    with TestClient(app) as test_client:
        yield {
            "client": test_client,
            "engine": engine,
            "session_factory": TestingSessionLocal,
            "upload_dir": test_upload_dir,
            "db_path": test_db_file,
            "vector_store": isolated_vector_store,
        }

    # Teardown & restore
    app.dependency_overrides.clear()
    settings.database_url = original_db_url
    settings.upload_dir = original_upload_dir
    session_module._engine = old_engine
    session_module._session_factory = old_sf
